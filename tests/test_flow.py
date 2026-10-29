"""Сквозной сценарий MVP: от импорта слушателя до сертификата с QR и отчета заказчику."""
import json
from datetime import date, datetime, timedelta

from tests.base import AppTestCase, xlsx_file

from app.models import (Certificate, Cohort, Customer, Enrollment, Evidence, EvidenceFile, Level, Program, Session,
                        TimecodeComment, User, db)


class FullFlowTest(AppTestCase):
    def upload_video(self, client, size=3000):
        data = self.video_bytes(size)
        r = client.post("/api/uploads", json={"name": "lesson.mp4", "size": len(data), "type": "video/mp4"})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True))
        up = r.get_json()
        chunk = up["chunk"]
        received = 0
        # первая часть, затем "обрыв" и повторная отправка той же части (сервер вернет фактический объем)
        r = client.put(f"/api/uploads/{up['id']}?offset=0", data=data[:chunk])
        received = r.get_json()["received"]
        r = client.put(f"/api/uploads/{up['id']}?offset=0", data=data[:chunk])
        self.assertEqual(r.get_json()["received"], received)  # повтор не дублирует данные
        status = client.get(f"/api/uploads/{up['id']}").get_json()
        self.assertEqual(status["received"], received)
        file_id = None
        while received < len(data):
            r = client.put(f"/api/uploads/{up['id']}?offset={received}", data=data[received:received + chunk])
            body = r.get_json()
            received = body["received"]
            file_id = body.get("file_id") or file_id
        self.assertIsNotNone(file_id)
        return file_id

    def test_complete_level_and_get_certificate(self):
        admin = self.new_client("admin@academy.demo")
        level = Level.query.join(Program).filter(Program.slug == "fll-future-edition").first()
        cust = Customer.query.first()
        trainer = User.query.filter_by(email="trainer@academy.demo").first()

        # 1. Админ создает когорту
        r = admin.post("/admin/cohorts/new", data={
            "name": "Шымкент, тест", "level_id": level.id, "start_date": date.today().isoformat(),
            "trainer_id": trainer.id, "customer_id": cust.id, "language": "kk", "capacity": 20, "status": "active"})
        self.assertEqual(r.status_code, 302)
        cohort = Cohort.query.filter_by(name="Шымкент, тест").first()
        self.assertEqual(cohort.program_version, level.program.version)

        # 2. Интенсив
        start = datetime.now().replace(microsecond=0) - timedelta(hours=1)
        admin.post(f"/admin/cohorts/{cohort.id}", data={
            "action": "add_session", "title": "Интенсив", "kind": "intensive",
            "starts_at": start.strftime("%Y-%m-%dT%H:%M"), "ends_at": (start + timedelta(hours=6)).strftime("%Y-%m-%dT%H:%M"),
            "location": "Шымкент"})
        self.assertEqual(len(cohort.sessions), 1)

        # 3. Импорт слушателей из Excel
        f = xlsx_file([["ФИО", "Email", "Телефон", "Организация", "Город"],
                       ["Жанна Тестова", "zhanna@test.kz", "+77000000000", "Школа № 1", "Шымкент"],
                       ["Без почты", "", "", "", ""]])
        r = admin.post(f"/admin/cohorts/{cohort.id}/import", data={"file": (f, "list.xlsx"), "send": "1"},
                       content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200)
        page = r.get_data(as_text=True)
        self.assertIn("zhanna@test.kz", page)
        user = User.query.filter_by(email="zhanna@test.kz").first()
        self.assertIsNotNone(user.invite_token)
        self.assertEqual(user.locale, "kk")

        # 4. Учитель активирует аккаунт по приглашению
        teacher = self.new_client()
        r = teacher.post(f"/invite/{user.invite_token}", data={"password": "superpass1", "consent": "1"})
        self.assertEqual(r.status_code, 302)
        db.session.refresh(user)
        self.assertIsNone(user.invite_token)
        self.assertIsNotNone(user.consent_at)

        # 5. Уроки и тест
        enr = Enrollment.query.filter_by(cohort_id=cohort.id, user_id=user.id).first()
        for lesson in level.lessons:
            url = f"/my/cohort/{cohort.id}/lesson/{lesson.id}"
            self.assertEqual(teacher.get(url).status_code, 200)
            if lesson.kind == "quiz":
                wrong = {f"q{q.id}": str((q.correct[0] + 1) % len(q.options)) for q in lesson.questions}
                r = teacher.post(url, data=wrong)
                self.assertIn("0%", r.get_data(as_text=True))
                right = {f"q{q.id}": str(q.correct[0]) for q in lesson.questions}
                teacher.post(url, data=right)
            else:
                teacher.post(url)
        from app.services import progress
        st = progress.state(enr)
        self.assertEqual(st.lessons_done, st.lessons_total)
        self.assertEqual(enr.status, "active")  # еще нет задания и посещения

        # 6. Задание: видео частями + согласие
        a = level.assignments[0]
        file_id = self.upload_video(teacher)
        url = f"/my/cohort/{cohort.id}/assignment/{a.id}"
        r = teacher.post(url, data={"file_id": file_id, "text": "5 класс, 12 детей"})  # без согласия
        self.assertEqual(Evidence.query.count(), 0)
        teacher.post(url, data={"file_id": file_id, "text": "5 класс, 12 детей", "consent_children": "1"})
        ev = Evidence.query.filter_by(enrollment_id=enr.id).first()
        self.assertEqual(ev.status, "submitted")
        self.assertEqual(len(ev.files), 1)

        # чужой слушатель не видит видео, тренер (он же проверяющий) видит
        other = self.new_client("teacher1@academy.demo")
        self.assertEqual(other.get(f"/files/{file_id}").status_code, 404)
        tr = self.new_client("trainer@academy.demo")
        r = tr.get(f"/files/{file_id}", headers={"Range": "bytes=0-99"})
        self.assertEqual(r.status_code, 206)

        # 7. Тренер оставляет комментарий по таймкоду и возвращает на доработку
        r = tr.post(f"/api/files/{file_id}/comments", json={"text": "Здесь задайте открытый вопрос", "seconds": 75.5})
        self.assertTrue(r.get_json()["ok"])
        self.assertEqual(TimecodeComment.query.count(), 1)
        criteria = {f"c{c.id}": "1" for c in a.criteria}
        r = tr.post(f"/review/{ev.id}", data={**criteria, "decision": "accepted"})  # ниже минимума
        db.session.refresh(ev)
        self.assertEqual(ev.status, "submitted")
        tr.post(f"/review/{ev.id}", data={**criteria, "decision": "returned", "summary": "Добавьте рефлексию"})
        db.session.refresh(ev)
        self.assertEqual(ev.status, "returned")
        self.assertIn("1:15", teacher.get(url).get_data(as_text=True))  # учитель видит таймкод

        # 8. Повторная сдача и принятие
        file2 = self.upload_video(teacher, 2500)
        teacher.post(url, data={"file_id": file2, "consent_children": "1"})
        ev2 = Evidence.query.filter_by(enrollment_id=enr.id, attempt=2).first()
        self.assertIsNotNone(ev2)
        good = {f"c{c.id}": "3" for c in a.criteria}
        tr.post(f"/review/{ev2.id}", data={**good, "decision": "accepted", "summary": "Отлично"})
        db.session.refresh(enr)
        self.assertEqual(enr.status, "active")  # не хватает интенсива

        # 9. Посещаемость (как из офлайн-очереди) завершает уровень
        session = cohort.sessions[0]
        r = tr.post(f"/api/sessions/{session.id}/attendance", json={"marks": [{"enrollment_id": enr.id, "present": True}]})
        self.assertEqual(r.get_json()["saved"], [enr.id])
        db.session.refresh(enr)
        self.assertEqual(enr.status, "completed")
        cert = Certificate.query.filter_by(enrollment_id=enr.id).first()
        self.assertIsNotNone(cert)
        self.assertEqual(cert.recipient_name, "Жанна Тестова")
        self.assertGreaterEqual(cert.hours, 6)

        # 10. Сертификат: PDF, проверка по QR, бейдж
        r = teacher.get(f"/my/certificates/{cert.code}.pdf")
        self.assertEqual(r.mimetype, "application/pdf")
        self.assertTrue(r.data.startswith(b"%PDF"))
        anon = self.new_client()
        r = anon.get(f"/verify/{cert.code}")
        self.assertIn("Жанна Тестова", r.get_data(as_text=True))
        badge = anon.get(f"/badges/{cert.code}.json").get_json()
        self.assertEqual(badge["credentialSubject"]["name"], "Жанна Тестова")
        self.assertEqual(anon.get("/verify/NOPE").status_code, 404)

        # 11. Отчет заказчика и Excel
        cust_client = self.new_client("customer@academy.demo")
        r = cust_client.get(f"/reports/cohort/{cohort.id}")
        self.assertIn("Жанна Тестова", r.get_data(as_text=True))
        self.assertNotIn("zhanna@test.kz", r.get_data(as_text=True))
        r = cust_client.get(f"/reports/cohort/{cohort.id}.xlsx")
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(r.data)).active
        flat = [c for row in ws.iter_rows(values_only=True) for c in row if c]
        self.assertIn("Жанна Тестова", flat)
        self.assertNotIn("zhanna@test.kz", flat)

        # 12. Отзыв сертификата
        admin.post("/admin/certificates", data={"action": "revoke", "certificate_id": cert.id})
        self.assertIn("отозван", anon.get(f"/verify/{cert.code}?lang=ru").get_data(as_text=True).lower())

    def test_customer_sees_only_own_cohorts(self):
        other = Customer(name="Другой заказчик")
        db.session.add(other)
        db.session.flush()
        level = Level.query.join(Program).filter(Program.slug == "fll-future-edition").first()
        c = Cohort(name="Чужая когорта", program_id=level.program_id, level_id=level.id, customer_id=other.id,
                   start_date=date.today())
        db.session.add(c)
        db.session.commit()
        client = self.new_client("customer@academy.demo")
        self.assertNotIn("Чужая когорта", client.get("/reports").get_data(as_text=True))
        self.assertEqual(client.get(f"/reports/cohort/{c.id}").status_code, 404)

    def test_open_enrollment_link(self):
        cohort = Cohort.query.first()
        anon = self.new_client()
        r = anon.get(f"/join/{cohort.invite_code}")
        self.assertEqual(r.status_code, 200)
        r = anon.post("/register", data={"full_name": "Новый Учитель", "email": "new@test.kz",
                                         "password": "password1", "consent": "1", "phone": "+77011112233"})
        self.assertEqual(r.status_code, 302)
        u = User.query.filter_by(email="new@test.kz").first()
        self.assertIsNotNone(Enrollment.query.filter_by(cohort_id=cohort.id, user_id=u.id).first())
        cohort.is_open = False
        db.session.commit()
        self.assertEqual(self.new_client().get(f"/join/{cohort.invite_code}").status_code, 410)

    def test_studio_builds_program(self):
        m = self.new_client("methodist@academy.demo")
        r = m.post("/studio", data={"title_ru": "Тестовая программа", "title_kk": "", "title_en": "Test program",
                                    "direction": "ftc"})
        p = Program.query.filter_by(slug="test-program").first()
        self.assertIsNotNone(p)
        lv = p.levels[0]
        m.post(f"/studio/levels/{lv.id}", data={"action": "add_module", "module_title_ru": "Модуль 1"})
        mod = lv.modules[0]
        r = m.post(f"/studio/levels/{lv.id}", data={"action": "add_lesson", "module_id": mod.id, "kind": "quiz"})
        lesson = mod.lessons[0]
        m.post(f"/studio/lessons/{lesson.id}", data={"action": "add_question"})
        q = lesson.questions[0]
        m.post(f"/studio/lessons/{lesson.id}", data={
            "action": "save_question", "question_id": q.id, "prompt_ru": "2+2?",
            "opt_0_ru": "3", "opt_1_ru": "", "opt_2_ru": "4", "correct": "2"})
        db.session.refresh(q)
        self.assertEqual(len(q.options), 2)
        self.assertEqual(q.correct, [1])  # индекс пересчитан после пустой строки
        m.post(f"/studio/levels/{lv.id}", data={"action": "add_assignment"})
        self.assertEqual(len(lv.assignments), 1)
        m.post(f"/studio/programs/{p.id}", data={"action": "publish"})
        db.session.refresh(p)
        self.assertEqual((p.status, p.version), ("published", 1))
        self.assertEqual(self.new_client().get(f"/programs/{p.slug}").status_code, 200)
        r = m.post("/studio/preview-md", json={"text": "**x** <script>alert(1)</script>"})
        html = r.get_json()["html"]
        self.assertIn("<strong>x</strong>", html)
        self.assertNotIn("<script>", html)

    def test_lead_form(self):
        r = self.client.post("/request", data={"org_name": "Школа", "contact_name": "Директор", "phone": "+7700"})
        self.assertEqual(r.status_code, 302)
        admin = self.new_client("admin@academy.demo")
        self.assertIn("Директор", admin.get("/admin/leads").get_data(as_text=True))

    def test_upload_rejects_bad_files(self):
        c = self.new_client("teacher1@academy.demo")
        self.assertEqual(c.post("/api/uploads", json={"name": "virus.exe", "size": 10}).status_code, 400)
        self.assertEqual(c.post("/api/uploads", json={"name": "v.mp4", "size": 10 ** 12}).status_code, 400)


class SecurityRegressionTest(AppTestCase):
    """Проверки по итогам ревью кода."""

    def test_register_cannot_take_over_invited_account(self):
        admin = self.new_client("admin@academy.demo")
        admin.post("/admin/users/new", data={"email": "boss@x.kz", "full_name": "Новый Админ", "roles": "admin"})
        anon = self.new_client()
        r = anon.post("/register", data={"full_name": "Злоумышленник", "email": "boss@x.kz",
                                         "password": "hacked123", "consent": "1"})
        self.assertEqual(r.status_code, 400)
        self.assertEqual(anon.get("/admin").status_code, 302)
        self.assertFalse(User.query.filter_by(email="boss@x.kz").first().check_password("hacked123"))

    def test_uploaded_html_is_never_served_inline(self):
        c = self.new_client("teacher1@academy.demo")
        r = c.post("/api/uploads", json={"name": "notes.txt", "size": 20, "type": "text/html"})
        up = r.get_json()
        fid = c.put(f"/api/uploads/{up['id']}?offset=0", data=b"<script>alert(1)</script>"[:20]).get_json()["file_id"]
        r = c.get(f"/files/{fid}")
        self.assertNotIn("text/html", r.headers["Content-Type"])
        self.assertIn("attachment", r.headers["Content-Disposition"])
        self.assertIn("sandbox", r.headers["Content-Security-Policy"])

    def test_capacity_respected_when_joining_after_login(self):
        cohort = Cohort.query.first()
        cohort.capacity = Enrollment.query.filter_by(cohort_id=cohort.id).count()
        db.session.commit()
        c = self.new_client()
        c.get(f"/join/{cohort.invite_code}")
        self.login("mentor@academy.demo", client=c)
        mentor = User.query.filter_by(email="mentor@academy.demo").first()
        self.assertIsNone(Enrollment.query.filter_by(cohort_id=cohort.id, user_id=mentor.id).first())

    def test_no_completion_without_intensive(self):
        from app.services import progress
        cohort = Cohort.query.first()
        for s in list(cohort.sessions):
            db.session.delete(s)
        db.session.commit()
        enr = cohort.enrollments[0]
        self.assertEqual(progress.state(enr).sessions_required, 1)

    def test_assignment_with_submissions_cannot_be_deleted(self):
        level = Level.query.join(Program).filter(Program.slug == "fll-future-edition").first()
        a = level.assignments[0]
        enr = Enrollment.query.first()
        db.session.add(Evidence(enrollment_id=enr.id, assignment_id=a.id))
        db.session.commit()
        m = self.new_client("methodist@academy.demo")
        m.post(f"/studio/assignments/{a.id}", data={"action": "delete"})
        self.assertIsNotNone(db.session.get(type(a), a.id))

    def test_excel_formula_is_neutralized(self):
        from io import BytesIO
        from openpyxl import load_workbook
        u = User.query.filter_by(email="teacher1@academy.demo").first()
        u.full_name = "=HYPERLINK(\"http://x\")"
        db.session.commit()
        cohort = Cohort.query.first()
        r = self.new_client("customer@academy.demo").get(f"/reports/cohort/{cohort.id}.xlsx")
        ws = load_workbook(BytesIO(r.data)).active
        types = {c.data_type for row in ws.iter_rows() for c in row if c.value and str(c.value).startswith(("'=", "="))}
        self.assertNotIn("f", types)

    def test_cyrillic_search_and_bad_comment_seconds(self):
        admin = self.new_client("admin@academy.demo")
        self.assertIn("Асель", admin.get("/admin/users?q=асель").get_data(as_text=True))
        self.assertEqual(admin.post("/admin/customers", data={"customer_id": 99999, "name": "x"}).status_code, 404)
