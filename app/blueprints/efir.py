"""Прямой эфир: видео с YouTube + анкета, викторина «как Kahoot», вопросы спикеру и счетчики в Академии.

Страницы:
- /efir/<slug>            учитель (видео + панель); без входа - приглашение зарегистрироваться
- /efir/<slug>/screen     экран для OBS (Browser Source, 1920x1080), доступ по ключу ?key=...
- /efir/<slug>/overlay    маленький счетчик и QR в углу поверх камеры спикера (прозрачный фон)
- /efir/<slug>/control    пульт оператора (методист, тренер, админ)
"""
from __future__ import annotations

import io
import re
import hashlib
import hmac
import secrets
from datetime import datetime

import segno
from flask import Blueprint, abort, current_app, g, jsonify, redirect, render_template, request, send_file, url_for
from markupsafe import Markup

from ..core import audit, role_required
from ..i18n import t
from ..models import EfirAnswer, EfirAsk, EfirEvent, EfirFeedback, EfirQuestion, EfirViewer, User, db, utcnow
from ..services import efir as svc
from ..services import efir_builder as lb
from ..services import efir_deck as deck_svc
from ..services import efir_media as media
from ..services import efir_video as vid_svc
from ..services.mail import absolute

bp = Blueprint("efir", __name__, url_prefix="/efir")

YT_RE = re.compile(r"(?:v=|youtu\.be/|/live/|/embed/|/shorts/)([A-Za-z0-9_-]{11})")


def _event(slug: str) -> EfirEvent:
    ev = EfirEvent.query.filter_by(slug=slug).first()
    if not ev:
        abort(404)
    return ev


def _screen_ok(ev: EfirEvent) -> bool:
    """Смотреть экран эфира: по ключу экрана, по ключу ссылки спикера (предпросмотр у спикера) или команде."""
    key = request.args.get("key") or ""
    if key and (secrets.compare_digest(key, ev.screen_key) or secrets.compare_digest(key, ev.guest_key)):
        return True
    return bool(g.user and g.user.has_role("methodist", "trainer"))


def qr_svg(url: str, dark: str = "#0b0b12", light: str = "#ffffff") -> Markup:
    return Markup(segno.make(url, error="m").svg_inline(scale=10, border=2, dark=dark, light=light, omitsize=True))


def _links(ev: EfirEvent) -> dict:
    links = dict(ev.links or {})
    links["academy"] = absolute(url_for("efir.watch", slug=ev.slug, lang=ev.lang or "kk"))
    return links


def _pretty(url: str) -> str:
    return re.sub(r"^https?://", "", url).split("?")[0].rstrip("/")


# ─── Учитель ─────────────────────────────────────────────────────────────
@bp.route("/<slug>", strict_slashes=False)
def watch(slug):
    ev = _event(slug)
    if not g.user:
        nxt = url_for("efir.watch", slug=slug)
        return render_template("efir/gate.html", ev=ev, nxt=nxt)
    v = svc.viewer_for(ev, g.user)
    links = _links(ev)
    return render_template("efir/watch.html", ev=ev, v=v, regions=svc.KZ_REGIONS, subjects=svc.SUBJECTS,
                           positions=svc.POSITIONS, robotics=svc.ROBOTICS, links=links,
                           bot_handle=lb.bot_handle(links.get("bot")), finish_text=_finish(ev),
                           fb=EfirFeedback.query.filter_by(event_id=ev.id, user_id=g.user.id).first(),
                           fb_liked=svc.FB_LIKED, fb_improve=svc.FB_IMPROVE)


@bp.route("/<slug>/state")
def state(slug):
    st = svc.public_state(slug)
    if st is None:
        abort(404)
    out = dict(st)
    out["yt"] = _yt_id(slug)
    if g.user:
        ev_id = _event_id(slug)
        v = EfirViewer.query.filter_by(event_id=ev_id, user_id=g.user.id).first()
        if v:
            svc.heartbeat(v)
        me = {"score": v.score if v else 0, "profile": bool(v and v.profile_at)}
        q = st.get("q")
        if q:
            a = EfirAnswer.query.filter_by(question_id=q["id"], user_id=g.user.id).first()
            if a:
                me["answer"] = {"choice": a.choice}
                if st["mode"] == "reveal":
                    me["answer"].update({"correct": a.is_correct, "points": a.points})
        if st["mode"] in ("leaderboard", "finish") and v:
            ev = db.session.get(EfirEvent, ev_id)
            me["rank"], me["total"] = svc.rank_of(ev, g.user)
        out["me"] = me
        if st["mode"] in ("leaderboard", "finish"):
            s = svc.stats(slug) or {}
            out["top"] = s.get("top", [])[:5]
    resp = jsonify(out)
    resp.headers["Cache-Control"] = "no-store"
    return resp


_ids: dict[str, tuple[int, str | None]] = {}


def _event_id(slug: str) -> int:
    if slug not in _ids:
        ev = _event(slug)
        _ids[slug] = (ev.id, ev.youtube_id)
    return _ids[slug][0]


def _yt_id(slug: str) -> str | None:
    import time
    key = "yt:" + slug
    hit = svc._cache.get(key)
    if hit and time.monotonic() - hit[0] < 5:
        return hit[1]
    ev = EfirEvent.query.filter_by(slug=slug).first()
    val = ev.youtube_id if ev else None
    svc._cache[key] = (time.monotonic(), val)
    return val


def _need_user():
    if not g.user:
        abort(401)


@bp.route("/<slug>/profile", methods=["POST"])
def profile(slug):
    _need_user()
    ev = _event(slug)
    f = request.get_json(silent=True) or request.form
    region = (f.get("region") or "").strip()
    school = (f.get("school") or "").strip()[:255]
    subjects = [s for s in (f.get("subjects") or []) if s in svc.SUBJECTS] if isinstance(f.get("subjects"), list) \
        else [s for s in request.form.getlist("subjects") if s in svc.SUBJECTS]
    if region not in svc.REGION_KEYS or len(school) < 2 or not subjects:
        return jsonify({"ok": False, "error": t("live.err_profile")}), 400
    v = svc.viewer_for(ev, g.user)
    v.region, v.school, v.subjects = region, school, subjects
    v.city = (f.get("city") or "").strip()[:128] or v.city
    pos = f.get("position")
    v.position = pos if pos in svc.POSITIONS else None
    rb = f.get("robotics")
    v.robotics = rb if rb in svc.ROBOTICS else None
    v.profile_at = v.profile_at or utcnow()
    # заполняем профиль Академии, если он пуст: школа и город пригодятся в курсах и отчетах
    if not g.user.organization:
        g.user.organization = school
    if not g.user.city and v.city:
        g.user.city = v.city
    if not g.user.region:
        g.user.region = svc.region_name(region, "ru")
    db.session.commit()
    return jsonify({"ok": True})


@bp.route("/<slug>/answer", methods=["POST"])
def answer(slug):
    _need_user()
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    q = db.session.get(EfirQuestion, int(f.get("qid") or 0))
    if not q or q.event_id != ev.id:
        abort(404)
    try:
        choice = int(f.get("choice"))
        ms = int(f["ms"]) if f.get("ms") is not None else None
    except (TypeError, ValueError):
        abort(400)
    return jsonify(svc.answer(ev, q, g.user, choice, ms))


@bp.route("/<slug>/ask", methods=["POST"])
def ask(slug):
    _need_user()
    ev = _event(slug)
    f = request.get_json(silent=True) or request.form
    ok, err, wait = svc.add_ask(ev, g.user, f.get("text") or "")
    if not ok:
        return jsonify({"ok": False, "error": t(err, n=wait), "wait": wait}), 429 if err in ("live.ask_wait", "live.ask_limit") else 400
    return jsonify({"ok": True, "cooldown": svc.ASK_COOLDOWN})


@bp.route("/<slug>/feedback", methods=["POST"])
def feedback(slug):
    _need_user()
    ev = _event(slug)
    fb, err = svc.save_feedback(ev, g.user, request.get_json(silent=True) or {})
    if err:
        return jsonify({"ok": False, "error": t(err)}), 400
    return jsonify({"ok": True})


# ─── Экран для OBS ───────────────────────────────────────────────────────
@bp.route("/<slug>/screen")
def screen(slug):
    ev = _event(slug)
    if not _screen_ok(ev):
        abort(403)
    links = _links(ev)
    qrs = {k: qr_svg(links[k]) for k in ("academy", "connect", "bot") if links.get(k)}
    from ..i18n import current_locale as get_locale
    return render_template("efir/screen.html", ev=ev, links=links, qrs=qrs, pretty=_pretty,
                           key=request.args.get("key", ""), subjects=svc.SUBJECTS, blocks_json=svc.blocks(slug),
                           starts=ev.starts_at.strftime("%Y-%m-%dT%H:%M:%S") if ev.starts_at else "",
                           kicker=lb.kicker(ev, get_locale()), finish_text=_finish(ev) or t("live.scr_finish_sub_def"),
                           bot_handle=lb.bot_handle(links.get("bot")), demo=_demo(ev), media_on=media.enabled(),
                           deck=deck_svc.for_screen(ev, links, request.args.get("key", "")))


def _finish(ev: EfirEvent) -> str:
    from ..i18n import current_locale as get_locale
    f = (ev.texts or {}).get("finish") or {}
    return f.get(get_locale()) or f.get("kk") or f.get("ru") or ""


def _demo(ev: EfirEvent) -> dict | None:
    """Предпросмотр сцены в конструкторе (?demo=waiting|talk|profile|question|reveal|leaderboard|finish): без эфира,
    с примерными цифрами. Только для команды: экран с ключом показывает живой эфир."""
    mode = request.args.get("demo")
    if mode not in svc.MODES or not (g.user and g.user.has_role("methodist", "trainer")):
        return None
    import time as _t
    qs = EfirQuestion.query.filter_by(event_id=ev.id).order_by(EfirQuestion.position).all()
    st = {"mode": mode, "v": 0, "now": _t.time()}
    if mode == "slides":
        try:
            st["slide"] = int(request.args.get("slide") or 0)
        except ValueError:
            st["slide"] = 0
    if mode in ("question", "reveal"):
        q = next((x for x in qs if str(x.id) == request.args.get("qid")), qs[0] if qs else None)
        if not q:
            st["mode"] = mode = "talk"
        else:
            st["q"] = {"id": q.id, "n": qs.index(q) + 1, "of": len(qs), "text": q.text, "options": q.options,
                       "seconds": q.seconds, "at": 0, "block": q.block}
            if mode == "reveal":
                st["q"].update(correct=q.correct, note=q.note)
    by = [64, 212, 41, 23][:len(st["q"]["options"])] if st.get("q") else [0, 0, 0, 0]
    top = [{"name": n, "school": sch, "region": r, "region_ru": r, "score": sc} for n, sch, r, sc in (
        ("Айгерим С.", "№12 мектеп", "Астана", 4820), ("Нурлан К.", "Келешек мектебі", "Алматы", 4610),
        ("Динара М.", "№5 гимназия", "Шымкент", 4390), ("Ерлан Т.", "НЗМ", "Ақтөбе", 4120), ("Сауле Б.", "№34 мектеп", "Павлодар", 3980))]
    stats = {"online": 312, "joined": 348, "profiles": 241, "schools": 186,
             "subjects": {"physics": 61, "informatics": 74, "math": 52, "technology": 38, "primary": 21},
             "regions": [{"kk": "Астана", "ru": "Астана", "n": 48}, {"kk": "Алматы", "ru": "Алматы", "n": 41},
                         {"kk": "Шымкент", "ru": "Шымкент", "n": 33}, {"kk": "Түркістан облысы", "ru": "Туркестанская область", "n": 27}],
             "answers": {"qid": st["q"]["id"] if st.get("q") else 0, "by": by, "total": sum(by)},
             "top": top, "region_top": [{"kk": "Астана", "ru": "Астана", "n": 48, "avg": 2410},
                                        {"kk": "Алматы", "ru": "Алматы", "n": 41, "avg": 2280}]}
    picked = {"text": "Командаға неше оқушы керек?", "name": "Айгерим С.", "school": "№12 мектеп"} if mode == "ask" else None
    return {"state": st, "stats": stats, "picked": picked}


@bp.route("/<slug>/overlay")
def overlay(slug):
    ev = _event(slug)
    if not _screen_ok(ev):
        abort(403)
    links = _links(ev)
    return render_template("efir/overlay.html", ev=ev, links=links, qr=qr_svg(links["academy"]), pretty=_pretty,
                           key=request.args.get("key", ""))


@bp.route("/<slug>/screen.json")
def screen_json(slug):
    ev = _event(slug)
    if not _screen_ok(ev):
        abort(403)
    out = {"state": svc.public_state(slug), "stats": svc.stats(slug)}
    st = out["state"] or {}
    if request.args.get("dv") is not None and str(st.get("dv", 0)) != request.args.get("dv"):
        out["deck"] = deck_svc.for_screen(ev, _links(ev), request.args.get("key", ""))
    out["stage"] = _stage_out(ev)
    if (st.get("mode") == "video") and st.get("video"):
        v = vid_svc.get(ev, st["video"].get("vid") or "")
        out["vnow"] = _video_src(ev, v, request.args.get("key", "")) if v else None
    picked = (EfirAsk.query.filter_by(event_id=ev.id, status="picked").order_by(EfirAsk.created_at.desc()).first())
    if picked:
        v = EfirViewer.query.filter_by(event_id=ev.id, user_id=picked.user_id).first()
        out["picked"] = {"text": picked.text, "name": svc.short_name(picked.user.full_name),
                         "school": (v.school if v else "") or ""}
    resp = jsonify(out)
    resp.headers["Cache-Control"] = "no-store"
    return resp


# ─── Пульт ───────────────────────────────────────────────────────────────
@bp.route("/<slug>/control")
@role_required("methodist", "trainer")
def control(slug):
    ev = _event(slug)
    qs = EfirQuestion.query.filter_by(event_id=ev.id).order_by(EfirQuestion.position).all()
    links = _links(ev)
    return render_template("efir/control.html", ev=ev, qs=qs, blocks=svc.blocks(slug), links=links,
                           qsec={str(q.id): q.seconds for q in qs},
                           qdata={str(q.id): {"block": q.block, "text": q.text, "options": q.options, "correct": q.correct,
                                              "note": q.note or {}, "seconds": q.seconds} for q in qs},
                           screen_url=absolute(url_for("efir.screen", slug=slug, key=ev.screen_key, lang="kk")),
                           overlay_url=absolute(url_for("efir.overlay", slug=slug, key=ev.screen_key, lang="kk")),
                           deck=deck_svc.for_screen(ev, links), media_on=media.enabled(), speaker_url=guest_url(ev),
                           stage=_stage_out(ev), has_stream_key=bool(ev.stream_key), videos=_videos_json(ev))


@bp.route("/<slug>/control.json")
@role_required("methodist", "trainer")
def control_json(slug):
    ev = _event(slug)
    from sqlalchemy import func
    per_q = dict(db.session.query(EfirAnswer.question_id, func.count(EfirAnswer.id))
                 .filter(EfirAnswer.event_id == ev.id).group_by(EfirAnswer.question_id).all())
    right = dict(db.session.query(EfirAnswer.question_id, func.count(EfirAnswer.id))
                 .filter(EfirAnswer.event_id == ev.id, EfirAnswer.is_correct.is_(True)).group_by(EfirAnswer.question_id).all())
    st = dict(ev.state or {})
    return jsonify({"state": st, "now": __import__("time").time(), "stats": svc.stats(slug),
                    "per_q": {str(k): [n, right.get(k, 0)] for k, n in per_q.items()}, "asks": svc.asks(ev)})


@bp.route("/<slug>/control/state", methods=["POST"])
@role_required("methodist", "trainer")
def control_state(slug):
    import time
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    mode = f.get("mode")
    if mode not in svc.MODES:
        abort(400)
    extra = {}
    if mode == "question":
        q = db.session.get(EfirQuestion, int(f.get("qid") or 0))
        if not q or q.event_id != ev.id:
            abort(400)
        extra = {"qid": q.id, "at": time.time(), "block": q.block}
    elif mode == "reveal":
        st = ev.state or {}
        if not st.get("qid"):
            abort(400)
        extra = {"qid": st["qid"], "at": st.get("at", 0), "block": st.get("block")}
    elif mode == "profile":
        extra = {}
    elif mode == "slides":
        n = len(deck_svc.get(ev)["slides"])
        if not n:
            abort(400)
        try:
            extra = {"slide": max(0, min(int(f.get("slide") or 0), n - 1))}
        except (TypeError, ValueError):
            abort(400)
        if (ev.state or {}).get("mode") == "slides":  # листание не засоряет журнал: пишем только вход в слайды
            svc.set_state(ev, mode, **extra)
            return jsonify({"ok": True, "state": ev.state})
    audit("live.state", "efir_event", ev.id, mode=mode, qid=extra.get("qid"))
    svc.set_state(ev, mode, **extra)
    return jsonify({"ok": True, "state": ev.state})


@bp.route("/<slug>/control/settings", methods=["POST"])
@role_required("methodist", "trainer")
def control_settings(slug):
    ev = _event(slug)
    f = request.get_json(silent=True) or request.form
    raw = (f.get("youtube") or "").strip()
    if raw:
        m = YT_RE.search(raw)
        yid = m.group(1) if m else (raw if re.fullmatch(r"[A-Za-z0-9_-]{11}", raw) else None)
        if not yid:
            return jsonify({"ok": False, "error": t("live.err_youtube")}), 400
        ev.youtube_id = yid
    elif "youtube" in f:
        ev.youtube_id = None
    starts = (f.get("starts_at") or "").strip()
    if starts:
        try:
            ev.starts_at = datetime.strptime(starts, "%Y-%m-%dT%H:%M")
        except ValueError:
            return jsonify({"ok": False, "error": "starts_at"}), 400
    audit("live.settings", "efir_event", ev.id, youtube_id=ev.youtube_id)
    db.session.commit()
    svc.drop_cache(ev.slug)
    svc._cache.pop("yt:" + ev.slug, None)
    _ids.pop(ev.slug, None)
    return jsonify({"ok": True, "youtube_id": ev.youtube_id})


@bp.route("/<slug>/control/mute/<int:user_id>", methods=["POST"])
@role_required("methodist", "trainer")
def control_mute(slug, user_id):
    ev = _event(slug)
    on = bool((request.get_json(silent=True) or {}).get("on", True))
    n = svc.mute(ev, user_id, on)
    audit("live.mute", "efir_event", ev.id, user_id=user_id, on=on, hidden=n)
    db.session.commit()
    return jsonify({"ok": True, "hidden": n})


@bp.route("/<slug>/control/ask/<int:ask_id>", methods=["POST"])
@role_required("methodist", "trainer")
def control_ask(slug, ask_id):
    ev = _event(slug)
    a = db.session.get(EfirAsk, ask_id)
    if not a or a.event_id != ev.id:
        abort(404)
    status = (request.get_json(silent=True) or {}).get("status")
    if status not in ("new", "picked", "done", "hidden"):
        abort(400)
    if status == "picked":
        EfirAsk.query.filter_by(event_id=ev.id, status="picked").update({"status": "done"})
    a.status = status
    db.session.commit()
    return jsonify({"ok": True})


def stats_key(ev: EfirEvent) -> str:
    """Ключ гостевой ссылки на статистику (для министерства и партнеров). Выводится из SECRET_KEY, без хранения."""
    secret = str(current_app.config.get("SECRET_KEY") or "").encode()
    return hmac.new(secret, f"efir-stats:{ev.slug}:{ev.id}".encode(), hashlib.sha256).hexdigest()[:24]


def _stats_access(ev: EfirEvent) -> str | None:
    """'staff' - команда USTEM, 'guest' - по ссылке с ключом (без персональных данных), None - нет доступа."""
    if g.user and g.user.has_role("methodist", "trainer"):
        return "staff"
    key = request.args.get("key") or ""
    if key and secrets.compare_digest(key, stats_key(ev)):
        return "guest"
    return None


def _stats_data(ev: EfirEvent, guest: bool) -> dict:
    from ..services import efir_stats
    data = efir_stats.compute(ev)
    if guest:  # гостям без ФИО и школ авторов вопросов
        for a in data["asks"]:
            a["name"] = ""
            a["school"] = ""
        for c in data["feedback"]["comments"]:
            c["name"] = ""
    return data


@bp.route("/<slug>/stats")
def stats_page(slug):
    """Статистика и аналитика эфира: команде USTEM и гостям по ссылке с ключом."""
    ev = _event(slug)
    access = _stats_access(ev)
    if not access:
        if not g.user:
            return redirect(url_for("auth.login", next=request.full_path))
        abort(403)
    guest = access == "guest"
    share_url = absolute(url_for("efir.stats_page", slug=slug, key=stats_key(ev), lang="ru")) if not guest else ""
    return render_template("efir/stats.html", ev=ev, data=_stats_data(ev, guest), blocks=svc.blocks(slug),
                           guest=guest, share_url=share_url, key=request.args.get("key", "") if guest else "")


@bp.route("/<slug>/stats.json")
def stats_json(slug):
    ev = _event(slug)
    access = _stats_access(ev)
    if not access:
        abort(403)
    resp = jsonify(_stats_data(ev, access == "guest"))
    resp.headers["Cache-Control"] = "no-store"
    return resp


@bp.route("/<slug>/control/question", methods=["POST"])
@bp.route("/<slug>/control/question/<int:qid>", methods=["POST"])
@role_required("methodist", "trainer")
def control_question(slug, qid=None):
    """Изменить вопрос викторины или добавить новый (без qid)."""
    ev = _event(slug)
    q = None
    if qid:
        q = db.session.get(EfirQuestion, qid)
        if not q or q.event_id != ev.id:
            abort(404)
    q, err = svc.save_question(ev, request.get_json(silent=True) or {}, q)
    if err:
        return jsonify({"ok": False, "error": t(err)}), 400
    audit("live.question", "efir_question", q.id, key=q.key)
    db.session.commit()
    return jsonify({"ok": True, "id": q.id})


@bp.route("/<slug>/control/question/<int:qid>/delete", methods=["POST"])
@role_required("methodist", "trainer")
def control_question_delete(slug, qid):
    ev = _event(slug)
    q = db.session.get(EfirQuestion, qid)
    if not q or q.event_id != ev.id:
        abort(404)
    if EfirAnswer.query.filter_by(question_id=q.id).first() or (ev.state or {}).get("qid") == q.id:
        return jsonify({"ok": False, "error": t("live.q_del_used")}), 400
    audit("live.question_delete", "efir_question", q.id, key=q.key)
    db.session.delete(q)
    db.session.commit()
    svc.drop_cache(ev.slug)
    return jsonify({"ok": True})


@bp.route("/<slug>/control/reset", methods=["POST"])
@role_required("admin")
def control_reset(slug):
    """Обнуление эфира после тестов. Только администратор и только с вводом кода эфира для подтверждения."""
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    if (f.get("confirm") or "").strip().lower() != ev.slug:
        return jsonify({"ok": False, "error": t("live.c_reset_bad")}), 400
    counts = svc.reset(ev)
    audit("live.reset", "efir_event", ev.id, **counts)
    db.session.commit()
    return jsonify({"ok": True, **counts})


@bp.route("/<slug>/export.xlsx")
@role_required("methodist", "trainer")
def export(slug):
    from openpyxl import Workbook
    ev = _event(slug)
    qs = EfirQuestion.query.filter_by(event_id=ev.id).order_by(EfirQuestion.position).all()
    answers: dict[tuple[int, int], EfirAnswer] = {(a.user_id, a.question_id): a
                                                  for a in EfirAnswer.query.filter_by(event_id=ev.id)}
    wb = Workbook()
    ws = wb.active
    ws.title = "Участники"
    head = ["ФИО", "Email", "Телефон", "Регион", "Город", "Школа", "Предметы", "Должность", "Опыт робототехники",
            "Анкета", "Впервые на эфире", "Минут на странице", "Баллы"] + [f"{q.key}" for q in qs]
    ws.append(head)
    rows = (db.session.query(EfirViewer, User).join(User, User.id == EfirViewer.user_id)
            .filter(EfirViewer.event_id == ev.id).order_by(EfirViewer.score.desc()).all())
    for v, u in rows:
        line = [u.full_name, u.email, u.phone or "", svc.region_name(v.region, "ru"), v.city or "", v.school or "",
                ", ".join(t(f"live.subj_{s}", locale="ru") for s in (v.subjects or [])),
                t(f"live.pos_{v.position}", locale="ru") if v.position else "",
                t(f"live.rob_{v.robotics}", locale="ru") if v.robotics else "",
                "да" if v.profile_at else "нет", svc.local_hm(v.first_seen), round((v.seconds or 0) / 60), v.score]
        for q in qs:
            a = answers.get((u.id, q.id))
            line.append("" if not a else ("+" if a.is_correct else "-"))
        ws.append(line)
    ws2 = wb.create_sheet("Вопросы спикеру")
    ws2.append(["Время", "ФИО", "Школа", "Вопрос", "Статус"])
    for a in svc.asks(ev, limit=10000):
        ws2.append([a["at"], a["name"], a["school"], a["text"], a["status"]])
    ws3 = wb.create_sheet("Отзывы")
    ws3.append(["Время", "ФИО", "Email", "Школа", "Оценка", "Понравилось", "Улучшить", "Комментарий", "Темы следующих эфиров"])
    vmap = {v.user_id: v for v, _ in rows}
    for fb in EfirFeedback.query.filter_by(event_id=ev.id).order_by(EfirFeedback.created_at):
        v = vmap.get(fb.user_id)
        ws3.append([svc.local_hm(fb.created_at), fb.user.full_name if fb.user else "", fb.user.email if fb.user else "",
                    (v.school if v else "") or "", fb.rating,
                    ", ".join(t(f"live.fbl_{x}", locale="ru") for x in fb.liked or []),
                    ", ".join(t(f"live.fbi_{x}", locale="ru") for x in fb.improve or []), fb.comment or "", fb.topics or ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"efir-{ev.slug}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ─── Конструктор эфиров ──────────────────────────────────────────────────
def _builder_json(ev: EfirEvent) -> dict:
    qs = EfirQuestion.query.filter_by(event_id=ev.id).order_by(EfirQuestion.position).all()
    used = {qid for (qid,) in db.session.query(EfirAnswer.question_id).filter(EfirAnswer.event_id == ev.id).distinct()}
    return {
        "slug": ev.slug, "title": ev.title or {}, "subtitle": ev.subtitle or {}, "lang": ev.lang or "kk",
        "starts_at": ev.starts_at.strftime("%Y-%m-%dT%H:%M") if ev.starts_at else "",
        "youtube": f"https://youtu.be/{ev.youtube_id}" if ev.youtube_id else "",
        "partners": (ev.texts or {}).get("partners", ""),
        "speaker_name": (ev.texts or {}).get("speaker_name", ""), "speaker_role": (ev.texts or {}).get("speaker_role", ""),
        "has_photo": deck_svc.speaker_path(ev).exists(), "kicker": (ev.texts or {}).get("kicker") or {},
        "finish": (ev.texts or {}).get("finish") or {}, "links": ev.links or {}, "blocks": ev.blocks or [],
        "archived": bool(ev.archived), "has_viewers": lb.has_viewers(ev),
        "questions": [{"id": q.id, "block": q.block, "text": q.text, "options": q.options, "correct": q.correct,
                       "note": q.note or {}, "seconds": q.seconds, "used": q.id in used} for q in qs],
        "ready": lb.readiness(ev),
        "deck": _deck_json(ev),
        "media_on": media.enabled(), "stream_url": ev.stream_url or "", "has_stream_key": bool(ev.stream_key),
        "videos": _videos_json(ev),
        "stream_key_tail": (ev.stream_key or "")[-4:],
        "urls": {
            "speaker": guest_url(ev),
            "watch": absolute(url_for("efir.watch", slug=ev.slug, lang=ev.lang or "kk")),
            "screen": absolute(url_for("efir.screen", slug=ev.slug, key=ev.screen_key, lang=ev.lang or "kk")),
            "overlay": absolute(url_for("efir.overlay", slug=ev.slug, key=ev.screen_key, lang=ev.lang or "kk")),
            "control": absolute(url_for("efir.control", slug=ev.slug)),
            "stats": absolute(url_for("efir.stats_page", slug=ev.slug)),
            "guest": absolute(url_for("efir.stats_page", slug=ev.slug, key=stats_key(ev), lang="ru")),
            "qr": url_for("efir.qr_file", slug=ev.slug),
        },
    }


def _deck_json(ev: EfirEvent) -> dict:
    d = deck_svc.get(ev)
    return {"cam": d["cam"], "camrect": deck_svc.CAM[d["cam"]],
            "slides": [dict(s, thumb=url_for("efir.slide_img", slug=ev.slug, name=f"{s['id']}-t.jpg")) for s in d["slides"]]}


def _bad(key: str, code: int = 400):
    return jsonify({"ok": False, "error": t(key)}), code


@bp.route("/", strict_slashes=False)
@role_required("methodist", "trainer")
def hub():
    """Все эфиры: создать, скопировать, открыть конструктор, пульт, аналитику."""
    from sqlalchemy import func
    evs = EfirEvent.query.order_by(EfirEvent.starts_at.desc().nullsfirst(), EfirEvent.id.desc()).all()
    counts = dict(db.session.query(EfirViewer.event_id, func.count(EfirViewer.id)).group_by(EfirViewer.event_id).all())
    qn = dict(db.session.query(EfirQuestion.event_id, func.count(EfirQuestion.id)).group_by(EfirQuestion.event_id).all())
    from ..services.efir_stats import _local
    now = _local(utcnow())
    rows = [{"ev": e, "status": lb.status(e, now), "viewers": counts.get(e.id, 0), "questions": qn.get(e.id, 0)} for e in evs]
    from ..models import LiveEvent
    lives = LiveEvent.query.filter_by(archived=False).order_by(LiveEvent.starts_at.desc().nullsfirst()).all()
    from .live import links_ctx
    return render_template("efir/hub.html", rows=rows, ai=lb.ai_enabled(), lives=lives, **links_ctx())


@bp.route("/import", methods=["POST"])
@role_required("methodist", "trainer")
def import_live():
    """Перенести вебинар из /live в студию (копия, /live не меняется)."""
    from ..models import LiveEvent
    f = request.get_json(silent=True) or request.form
    src = LiveEvent.query.filter_by(slug=f.get("slug") or "").first()
    if not src:
        return _bad("efir.import_none", 404)
    ev = lb.import_live(src, g.user.id)
    audit("live.create", "efir_event", ev.id, slug=ev.slug, from_live=src.slug)
    db.session.commit()
    return jsonify({"ok": True, "slug": ev.slug, "url": url_for("efir.build", slug=ev.slug)})


@bp.route("/new", methods=["POST"])
@role_required("methodist", "trainer")
def create():
    f = request.get_json(silent=True) or request.form
    src = None
    if f.get("copy"):
        src = EfirEvent.query.filter_by(slug=f.get("copy")).first()
    title = f.get("title")
    if isinstance(title, str):  # название на одном языке: второй допишут в конструкторе
        title = {lb.guess_lang(title): title}
    try:
        ev = lb.create(title or {}, f.get("starts_at"), g.user.id, source=src,
                       with_questions=str(f.get("with_questions", "1")) not in ("0", "false", ""))
    except lb.BuildError as e:
        return _bad(str(e))
    audit("live.create", "efir_event", ev.id, slug=ev.slug, copy=src.slug if src else None)
    db.session.commit()
    return jsonify({"ok": True, "slug": ev.slug, "url": url_for("efir.build", slug=ev.slug)})


@bp.route("/<slug>/build")
@role_required("methodist", "trainer")
def build(slug):
    ev = _event(slug)
    others = EfirEvent.query.filter(EfirEvent.id != ev.id).order_by(EfirEvent.id.desc()).all()
    return render_template("efir/build.html", ev=ev, data=_builder_json(ev), ai=lb.ai_enabled(), others=others)


@bp.route("/<slug>/build.json")
@role_required("methodist", "trainer")
def build_json(slug):
    resp = jsonify(_builder_json(_event(slug)))
    resp.headers["Cache-Control"] = "no-store"
    return resp


@bp.route("/<slug>/build/save", methods=["POST"])
@role_required("methodist", "trainer")
def build_save(slug):
    ev = _event(slug)
    try:
        changed = lb.save_settings(ev, request.get_json(silent=True) or {})
    except lb.BuildError as e:
        db.session.rollback()
        return _bad(str(e))
    if "slug" in changed:
        audit("live.slug", "efir_event", ev.id, slug=ev.slug)
        db.session.commit()
    return jsonify({"ok": True, "data": _builder_json(ev)})


@bp.route("/<slug>/build/reorder", methods=["POST"])
@role_required("methodist", "trainer")
def build_reorder(slug):
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    moves = f.get("blocks") or {}
    if moves:  # перенос вопроса в другой блок при перетаскивании
        for qid, b in moves.items():
            q = db.session.get(EfirQuestion, int(qid))
            if q and q.event_id == ev.id and 1 <= int(b) <= len(ev.blocks or []):
                q.block = int(b)
    lb.reorder(ev, [int(x) for x in f.get("ids") or []])
    return jsonify({"ok": True})


@bp.route("/<slug>/build/paste", methods=["POST"])
@role_required("methodist", "trainer")
def build_paste(slug):
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    items, blocks = lb.parse_paste(f.get("text") or "")
    if f.get("dry"):
        return jsonify({"ok": True, "items": items, "blocks": blocks})
    if not items:
        return _bad("live.b_paste_none")
    try:
        block = max(1, min(int(f.get("block") or 1), len(ev.blocks or [1])))
        seconds = max(10, min(int(f.get("seconds") or 30), 120))
    except (TypeError, ValueError):
        block, seconds = 1, 30
    n = lb.add_parsed(ev, items, blocks, f.get("lang") or "ru", block, seconds)
    audit("live.paste", "efir_event", ev.id, n=n)
    db.session.commit()
    return jsonify({"ok": True, "added": n, "data": _builder_json(ev)})


@bp.route("/<slug>/build/ai", methods=["POST"])
@role_required("methodist", "trainer")
def build_ai(slug):
    """ИИ составляет вопросы по презентации (PDF) или тексту. Возвращает черновики, в эфир их добавляет /build/add."""
    ev = _event(slug)
    if not lb.ai_enabled():
        return _bad("live.b_ai_off")
    pdf = None
    up = request.files.get("pdf")
    if up and up.filename:
        pdf = up.read()
        if not pdf.startswith(b"%PDF") or len(pdf) > 30 * 1024 * 1024:
            return _bad("live.b_ai_pdf")
    text = request.form.get("text") or ""
    if not pdf and len(text.strip()) < 80:
        return _bad("live.b_ai_need")
    try:
        count = max(1, min(int(request.form.get("count") or 5), 15))
    except ValueError:
        count = 5
    try:
        items = lb.ai_questions(ev, count, pdf, text, request.form.get("focus") or "")
    except Exception:
        import logging
        logging.getLogger("academy.live").exception("live ai questions failed")
        return _bad("live.b_ai_fail", 502)
    audit("live.ai_questions", "efir_event", ev.id, n=len(items))
    db.session.commit()
    return jsonify({"ok": True, "items": items})


@bp.route("/<slug>/build/add", methods=["POST"])
@role_required("methodist", "trainer")
def build_add(slug):
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    n = 0
    for it in (f.get("items") or [])[:30]:
        data = dict(it)
        data["block"] = f.get("block") or 1
        data["seconds"] = f.get("seconds") or 30
        q, err = svc.save_question(ev, data)
        n += 1 if q else 0
    return jsonify({"ok": True, "added": n, "data": _builder_json(ev)})


@bp.route("/<slug>/build/translate", methods=["POST"])
@role_required("methodist", "trainer")
def build_translate(slug):
    ev = _event(slug)
    if not lb.ai_enabled():
        return _bad("live.b_ai_off")
    try:
        n = lb.ai_translate(ev)
    except Exception:
        import logging
        logging.getLogger("academy.live").exception("live ai translate failed")
        db.session.rollback()
        return _bad("live.b_ai_fail", 502)
    audit("live.ai_translate", "efir_event", ev.id, n=n)
    db.session.commit()
    return jsonify({"ok": True, "filled": n, "data": _builder_json(ev)})


@bp.route("/<slug>/build/delete", methods=["POST"])
@role_required("methodist", "trainer")
def build_delete(slug):
    ev = _event(slug)
    if (request.get_json(silent=True) or {}).get("confirm") != ev.slug:
        return _bad("live.c_reset_bad")
    try:
        eid = ev.id
        lb.delete(ev)
    except lb.BuildError as e:
        return _bad(str(e))
    audit("live.delete", "efir_event", eid, slug=slug)
    db.session.commit()
    return jsonify({"ok": True, "url": url_for("efir.hub")})


@bp.route("/<slug>/qr.svg")
@role_required("methodist", "trainer")
def qr_file(slug):
    """QR на страницу эфира для афиш и рассылок."""
    from flask import Response
    ev = _event(slug)
    url = absolute(url_for("efir.watch", slug=ev.slug, lang=ev.lang or "kk"))
    svg = segno.make(url, error="m").svg_inline(scale=12, border=3, dark="#0b0b12", light="#ffffff")
    return Response('<?xml version="1.0" encoding="UTF-8"?>\n' + svg, mimetype="image/svg+xml",
                    headers={"Content-Disposition": f'attachment; filename="qr-{ev.slug}.svg"'})


# ─── Презентация ─────────────────────────────────────────────────────────
@bp.route("/<slug>/slides/<name>")
def slide_img(slug, name):
    """Картинка слайда: для экрана OBS (по ключу) и команды. Имена неизменяемые, кешируются надолго."""
    ev = _event(slug)
    if not _screen_ok(ev) or not re.fullmatch(r"[0-9a-f]{16}(-t)?\.jpg", name):
        abort(404)
    path = deck_svc.folder(ev) / name
    if not path.exists():
        abort(404)
    return send_file(path, mimetype="image/jpeg", max_age=30 * 24 * 3600)


@bp.route("/<slug>/build/deck/upload", methods=["POST"])
@role_required("methodist", "trainer")
def build_deck_upload(slug):
    ev = _event(slug)
    up = request.files.get("pdf")
    if not up or not up.filename:
        return _bad("live.d_err_pdf")
    try:
        n = deck_svc.render_pdf(ev, up.read(), replace=request.form.get("replace") == "1")
    except deck_svc.DeckError as e:
        return _bad(str(e))
    audit("live.deck", "efir_event", ev.id, added=n)
    db.session.commit()
    return jsonify({"ok": True, "added": n, "deck": _deck_json(ev)})


@bp.route("/<slug>/build/deck", methods=["POST"])
@role_required("methodist", "trainer")
def build_deck(slug):
    ev = _event(slug)
    deck_svc.save(ev, request.get_json(silent=True) or {})
    svc.drop_cache(ev.slug)
    st = ev.state or {}
    n = len(deck_svc.get(ev)["slides"])
    if st.get("mode") == "slides" and (not n or st.get("slide", 0) >= n):  # удалили показанный слайд
        svc.set_state(ev, "slides" if n else "talk", **({"slide": n - 1} if n else {}))
    return jsonify({"ok": True, "deck": _deck_json(ev)})


@bp.route("/<slug>/build/deck/export.<kind>")
@role_required("methodist", "trainer")
def build_deck_export(slug, kind):
    """Скачать презентацию с окном камеры и QR: PDF для просмотра или ZIP с PNG (прозрачное окно) для OBS."""
    from flask import Response
    ev = _event(slug)
    if kind not in ("pdf", "zip"):
        abort(404)
    labels = {"cam": t("live.d_cam_here"), "watch": t("live.d_qr_watch"), "connect": t("live.d_qr_connect"),
              "bot": t("live.d_qr_bot"), "url": ""}
    try:
        data, mime, name = deck_svc.export(ev, _links(ev), kind, labels)
    except deck_svc.DeckError as e:
        return _bad(str(e))
    return Response(data, mimetype=mime, headers={"Content-Disposition": f'attachment; filename="{name}"'})


@bp.route("/<slug>/youtube-cover.jpg")
@role_required("methodist", "trainer")
def youtube_cover(slug):
    """Обложка для YouTube 1280 x 720. ?lang=kk|ru, ?dl=1 - скачать файлом."""
    from flask import Response
    ev = _event(slug)
    lang = request.args.get("lang") if request.args.get("lang") in ("kk", "ru") else (ev.lang or "kk")
    labels = {"live": t("live.yt_live", locale=lang)}
    data = deck_svc.youtube_cover(ev, lang, labels, absolute(url_for("efir.watch", slug=ev.slug)))
    headers = {"Cache-Control": "no-store"}
    if request.args.get("dl"):
        headers["Content-Disposition"] = f'attachment; filename="youtube-{ev.slug}-{lang}.jpg"'
    return Response(data, mimetype="image/jpeg", headers=headers)


@bp.route("/<slug>/build/speaker", methods=["POST"])
@role_required("methodist", "trainer")
def build_speaker(slug):
    ev = _event(slug)
    if request.form.get("remove"):
        deck_svc.speaker_path(ev).unlink(missing_ok=True)
        return jsonify({"ok": True})
    up = request.files.get("photo")
    if not up or not up.filename:
        return _bad("live.yt_err_photo")
    try:
        deck_svc.save_speaker(ev, up.read())
    except deck_svc.DeckError as e:
        return _bad(str(e))
    return jsonify({"ok": True})


# ─── Студия: видеосервер, спикеры по ссылке, трансляция в YouTube ────────


def _guest_ok(ev: EfirEvent) -> bool:
    k = request.args.get("k") or (request.get_json(silent=True) or {}).get("k") or ""
    return bool(k) and secrets.compare_digest(str(k), ev.guest_key)


def guest_url(ev: EfirEvent) -> str:
    return absolute(url_for("efir.guest", slug=ev.slug, k=ev.guest_key))


def _stage_out(ev: EfirEvent) -> dict:
    s = media.stage(ev)
    return {"on": s["on"], "muted": s["muted"], "hands": s["hands"], "layout": s["layout"], "v": s["v"]}


@bp.route("/<slug>/guest")
def guest(slug):
    """Страница спикера: проверка камеры и микрофона, имя, вход в студию. Вход не нужен, доступ по ключу ссылки."""
    ev = _event(slug)
    if not _guest_ok(ev):
        abort(404)
    deck = deck_svc.get(ev)
    thumbs = [url_for("efir.slide_img", slug=ev.slug, name=f"{s['id']}-t.jpg", key=ev.guest_key) for s in deck["slides"]]
    return render_template("efir/guest.html", ev=ev, k=ev.guest_key, media_on=media.enabled(),
                           lk_url=media.cfg()["url"], thumbs=thumbs,
                           preview_url=url_for("efir.screen", slug=ev.slug, key=ev.guest_key, preview=1,
                                               lang=ev.lang or "kk"))


@bp.route("/<slug>/guest/token", methods=["POST"])
def guest_token(slug):
    ev = _event(slug)
    if not _guest_ok(ev):
        abort(404)
    if not media.enabled():
        return _bad("efir.m_off", 503)
    f = request.get_json(silent=True) or {}
    name = " ".join(str(f.get("name") or "").split())[:80]
    role = " ".join(str(f.get("role") or "").split())[:80]
    if len(name) < 2:
        return _bad("efir.g_err_name")
    ident = media.guest_identity(f.get("prev"))
    tok = media.token(ev, ident, name=name, publish=True, meta={"role": role})
    audit("live.guest", "efir_event", ev.id, identity=ident, name=name)
    db.session.commit()
    return jsonify({"ok": True, "url": media.cfg()["url"], "token": tok, "identity": ident})


@bp.route("/<slug>/guest/status")
def guest_status(slug):
    ev = _event(slug)
    if not _guest_ok(ev):
        abort(404)
    s = media.stage(ev)
    ident = request.args.get("id") or ""
    st = ev.state or {}
    resp = jsonify({"on": ident in s["on"], "muted": ident in s["muted"], "hand": ident in s["hands"], "on_list": s["on"],
                    "slide": st.get("slide", 0) if st.get("mode") == "slides" else None, "mode": (ev.state or {}).get("mode") or "waiting",
                    "streaming": (ev.egress or {}).get("status") in media.ACTIVE})
    resp.headers["Cache-Control"] = "no-store"
    return resp


@bp.route("/<slug>/guest/slide", methods=["POST"])
def guest_slide(slug):
    """Спикер в эфире листает слайды сам. За кулисами нельзя: ведущий выводит спикера в эфир, и пульт видит то же самое."""
    ev = _event(slug)
    if not _guest_ok(ev):
        abort(404)
    f = request.get_json(silent=True) or {}
    ident = str(f.get("id") or "")[:24]
    if ident not in media.stage(ev)["on"]:
        return _bad("efir.g_slide_back", 403)
    n = len(deck_svc.get(ev)["slides"])
    if not n:
        return _bad("efir.m_bad")
    try:
        i = max(0, min(int(f.get("slide") or 0), n - 1))
    except (TypeError, ValueError):
        return _bad("efir.m_bad")
    if (ev.state or {}).get("mode") != "slides":
        audit("live.state", "efir_event", ev.id, mode="slides", by=ident)
    svc.set_state(ev, "slides", slide=i)
    return jsonify({"ok": True, "slide": i, "of": n})


@bp.route("/<slug>/guest/hand", methods=["POST"])
def guest_hand(slug):
    """Спикер поднимает или опускает руку: ведущий видит ✋ на пульте."""
    ev = _event(slug)
    if not _guest_ok(ev):
        abort(404)
    f = request.get_json(silent=True) or {}
    ident = str(f.get("id") or "")[:24]
    if not ident.startswith("g-"):
        return _bad("efir.m_bad")
    s = media.stage(ev)
    hands = [h for h in s["hands"] if h != ident] + ([ident] if f.get("up") and ident not in s["on"] else [])
    s = media.set_stage(ev, hands=hands)
    db.session.commit()
    return jsonify({"ok": True, "hand": ident in s["hands"]})


@bp.route("/<slug>/media.json")
def media_json(slug):
    """Невидимый вход в комнату для экрана эфира (по ключу экрана) и для пульта (команда)."""
    ev = _event(slug)
    if not _screen_ok(ev):
        abort(403)
    if not media.enabled():
        return jsonify({"ok": False})
    kind = "control" if request.args.get("as") == "control" and g.user and g.user.has_role("methodist", "trainer") else "screen"
    tok = media.token(ev, media.viewer_identity(kind), hidden=True)
    resp = jsonify({"ok": True, "url": media.cfg()["url"], "token": tok, "stage": _stage_out(ev)})
    resp.headers["Cache-Control"] = "no-store"
    return resp


@bp.route("/<slug>/control/stage", methods=["POST"])
@role_required("methodist", "trainer")
def control_stage(slug):
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    s = media.set_stage(ev, on=f.get("on"), muted=f.get("muted"), layout=f.get("layout"), hands=f.get("hands"))
    audit("live.stage", "efir_event", ev.id, on=s["on"], layout=s["layout"])
    db.session.commit()
    return jsonify({"ok": True, "stage": _stage_out(ev)})


@bp.route("/<slug>/control/kick", methods=["POST"])
@role_required("methodist", "trainer")
def control_kick(slug):
    ev = _event(slug)
    ident = str((request.get_json(silent=True) or {}).get("identity") or "")
    if not ident.startswith("g-"):
        return _bad("efir.m_bad")
    s = media.stage(ev)
    media.set_stage(ev, on=[x for x in s["on"] if x != ident])
    try:
        media.remove(ev, ident)
    except media.MediaError:
        pass
    audit("live.kick", "efir_event", ev.id, identity=ident)
    db.session.commit()
    return jsonify({"ok": True, "stage": _stage_out(ev)})


def _stream_out(ev: EfirEvent) -> dict:
    eg = media.stream_status(ev)
    db.session.commit()
    return {"status": eg.get("status") or "", "error": eg.get("error") or "", "at": eg.get("at") or 0,
            "has_key": bool(ev.stream_key), "media": media.enabled()}


@bp.route("/<slug>/control/stream", methods=["GET", "POST"])
@role_required("methodist", "trainer")
def control_stream(slug):
    ev = _event(slug)
    if request.method == "GET":
        return jsonify({"ok": True, "stream": _stream_out(ev)})
    action = (request.get_json(silent=True) or {}).get("action")
    try:
        if action == "start":
            url = absolute(url_for("efir.screen", slug=ev.slug, key=ev.screen_key, lang=ev.lang or "kk", egress=1))
            media.start_stream(ev, url)
        elif action == "stop":
            media.stop_stream(ev)
        else:
            return _bad("efir.m_bad")
    except media.MediaError as e:
        db.session.rollback()
        msg = str(e)
        return jsonify({"ok": False, "error": t(msg) if msg.startswith("efir.") else t("efir.s_err", msg=msg)}), 502
    audit("live.stream", "efir_event", ev.id, op=action)
    db.session.commit()
    return jsonify({"ok": True, "stream": _stream_out(ev)})



# ─── Сцена «Видео»: YouTube или файл играет на экране эфира ─────────────
def _video_src(ev: EfirEvent, v: dict, key: str = "") -> dict:
    out = {"id": v["id"], "kind": v["kind"], "title": v.get("title") or "", "dur": v.get("dur")}
    if v["kind"] == "youtube":
        out["yt"] = v["yt"]
    else:
        out["src"] = url_for("efir.video_file", slug=ev.slug, vid=v["id"], ext=v["ext"], key=key or None)
    return out


def _videos_json(ev: EfirEvent) -> list[dict]:
    return [dict(_video_src(ev, v), size=v.get("size")) for v in vid_svc.items(ev)]


@bp.route("/<slug>/video/<vid>.<ext>")
def video_file(slug, vid, ext):
    """Файл видео для экрана эфира (по ключу) и команды. Поддерживает перемотку (Range)."""
    ev = _event(slug)
    v = vid_svc.get(ev, vid)
    if not _screen_ok(ev) or not v or v["kind"] != "file" or v["ext"] != ext:
        abort(404)
    p = vid_svc.path(ev, v)
    if not p.exists():
        abort(404)
    mime = {"mp4": "video/mp4", "m4v": "video/mp4", "webm": "video/webm", "mov": "video/quicktime"}[ext]
    return send_file(p, mimetype=mime, conditional=True, max_age=3600)


@bp.route("/<slug>/build/video", methods=["POST"])
@role_required("methodist", "trainer")
def build_video(slug):
    """Добавить видео (ссылка YouTube или загруженный файл), переименовать или удалить."""
    from ..models import EvidenceFile
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    try:
        if f.get("delete"):
            vid_svc.remove(ev, str(f["delete"]))
        elif f.get("id"):
            vid_svc.update(ev, str(f["id"]), title=f.get("title"), dur=f.get("dur"))
        elif f.get("file_id"):
            vid_svc.add_file(ev, db.session.get(EvidenceFile, int(f["file_id"])), g.user.id, f.get("title"), f.get("dur"))
        else:
            vid_svc.add_youtube(ev, f.get("url") or "", f.get("title"), f.get("dur"))
    except (vid_svc.VideoError, ValueError, TypeError) as e:
        db.session.rollback()
        msg = str(e)
        return _bad(msg if msg.startswith("efir.") else "efir.v_err_file")
    audit("live.video", "efir_event", ev.id, n=len(vid_svc.items(ev)))
    db.session.commit()
    return jsonify({"ok": True, "videos": _videos_json(ev)})


@bp.route("/<slug>/control/video", methods=["POST"])
@role_required("methodist", "trainer")
def control_video(slug):
    """Пульт: play (vid, from), pause, resume, seek (pos), stop (вернуться на сцену back)."""
    import time as _t
    ev = _event(slug)
    f = request.get_json(silent=True) or {}
    act = f.get("action")
    st = dict(ev.state or {})
    now = _t.time()
    if act == "play":
        v = vid_svc.get(ev, str(f.get("vid") or ""))
        if not v:
            return _bad("efir.v_err_file", 404)
        back = st.get("mode") if st.get("mode") not in (None, "video") else st.get("back", "talk")
        if st.get("mode") != "video":
            audit("live.state", "efir_event", ev.id, mode="video")
        svc.set_state(ev, "video", vid=v["id"], pos=max(0.0, float(f.get("from") or 0)), play_at=now, paused=False,
                      back=back, slide=st.get("slide", 0))
    elif st.get("mode") != "video":
        return _bad("efir.m_bad")
    elif act == "pause":
        svc.set_state(ev, "video", **{k: v for k, v in dict(st, pos=vid_svc.position(st), play_at=now, paused=True).items()
                                      if k not in ("mode", "v")})
    elif act == "resume":
        svc.set_state(ev, "video", **{k: v for k, v in dict(st, play_at=now, paused=False).items() if k not in ("mode", "v")})
    elif act == "seek":  # pos - секунда ролика
        try:
            pos = max(0.0, float(f.get("pos")))
        except (TypeError, ValueError):
            return _bad("efir.m_bad")
        svc.set_state(ev, "video", **{k: v for k, v in dict(st, pos=pos, play_at=now).items() if k not in ("mode", "v")})
    elif act == "stop":
        back = st.get("back") if st.get("back") in svc.MODES and st.get("back") != "video" else "talk"
        extra = {"slide": st.get("slide", 0)} if back == "slides" else {}
        svc.set_state(ev, back, **extra)
    else:
        return _bad("efir.m_bad")
    st = dict(ev.state or {})
    return jsonify({"ok": True, "state": st, "now": _t.time()})
