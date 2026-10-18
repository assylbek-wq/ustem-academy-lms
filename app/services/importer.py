"""Импорт слушателей в когорту из Excel или CSV и выгрузка отчетов в Excel."""
from __future__ import annotations

import csv
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..models import Enrollment, User, db

HEADERS = {
    "full_name": ("фио", "фамилия имя", "аты-жөні", "аты жөні", "full name", "name", "имя"),
    "email": ("email", "e-mail", "почта", "электронная почта", "эл. почта"),
    "phone": ("телефон", "phone", "тел", "телефон нөмірі"),
    "organization": ("организация", "школа", "мектеп", "ұйым", "organization", "school"),
    "city": ("город", "қала", "city", "населенный пункт"),
}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def _map_headers(row) -> dict[str, int]:
    out = {}
    for i, h in enumerate(row):
        n = _norm(h)
        for key, aliases in HEADERS.items():
            if key not in out and n in aliases:
                out[key] = i
    return out


def read_rows(filename: str, data: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = [list(r) for r in wb.active.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        raw = list(csv.reader(io.StringIO(text), dialect))
    raw = [r for r in raw if any(str(c or "").strip() for c in r)]
    if not raw:
        return []
    cols = _map_headers(raw[0])
    if "full_name" not in cols or "email" not in cols:
        raise ValueError("headers")
    rows = []
    for r in raw[1:]:
        item = {k: str(r[i]).strip() if i < len(r) and r[i] is not None else "" for k, i in cols.items()}
        rows.append(item)
    return rows


def import_into_cohort(cohort, rows: list[dict]) -> dict:
    """Создает недостающих пользователей (без пароля, с приглашением) и записывает в когорту."""
    report = {"created": [], "enrolled": 0, "already": 0, "errors": []}
    for n, r in enumerate(rows, start=2):
        email = (r.get("email") or "").lower()
        name = r.get("full_name") or ""
        if not EMAIL_RE.match(email) or not name:
            report["errors"].append((n, email or name or "-"))
            continue
        user = User.query.filter_by(email=email).first()
        if not user:
            user = User(email=email, full_name=name, phone=r.get("phone") or None,
                        organization=r.get("organization") or None, city=r.get("city") or None,
                        locale=cohort.language or "ru")
            user.issue_invite()
            db.session.add(user)
            db.session.flush()
            report["created"].append(user)
        if Enrollment.query.filter_by(cohort_id=cohort.id, user_id=user.id).first():
            report["already"] += 1
            continue
        db.session.add(Enrollment(cohort_id=cohort.id, user_id=user.id))
        report["enrolled"] += 1
    db.session.flush()
    return report


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Слушатели"
    ws.append(["ФИО", "Email", "Телефон", "Организация", "Город"])
    ws.append(["Иванова Айгерим Серикқызы", "aigerim@example.kz", "+7 700 000 0000", "Школа-лицей №1", "Алматы"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDE", (36, 30, 18, 30, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    """Защита от формул в Excel: текст, начинающийся с = + - @, сохраняется как текст."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def cohort_report_xlsx(cohort, summary: dict, labels: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = labels["sheet"][:31]
    ws.append([_cell(cohort.name)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([labels["program"], labels["program_value"]])
    ws.append([labels["dates"], labels["dates_value"]])
    ws.append([labels["learners"], summary["count"], labels["completed"], summary["completed"],
               labels["avg"], f'{summary["avg_percent"]}%'])
    ws.append([])
    head = labels["columns"]
    ws.append(head)
    hdr_row = ws.max_row
    fill = PatternFill("solid", fgColor="E1EFFF")
    for c in ws[hdr_row]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for enr, st in summary["rows"]:
        ws.append([_cell(v) for v in [
            enr.user.full_name, *([enr.user.email] if labels.get("with_email") else []),
            enr.user.organization or "", enr.user.city or "",
            f"{st.lessons_done}/{st.lessons_total}", f"{st.assignments_accepted}/{st.assignments_total}",
            f"{st.sessions_attended}/{st.sessions_required}" if st.sessions_required else "-",
            f"{st.percent}%", labels["status"].get(enr.status, enr.status),
        ]])
    for col, w in zip("ABCDEFGHI", (34, 30, 28, 14, 12, 12, 14, 10, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
